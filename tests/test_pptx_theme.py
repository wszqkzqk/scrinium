from __future__ import annotations

import os
import shutil
import subprocess
import sys
from dataclasses import FrozenInstanceError
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest

pytest.importorskip("pptx")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scrinium.presentation.pptx_style import (  # noqa: E402
    base_slide,
    chapter_slide,
    configure_theme,
    new_presentation,
    rich_text,
    text,
)
from scrinium.presentation.theme import DEFAULT_THEME, replace  # noqa: E402


@pytest.fixture(autouse=True)
def reset_theme():
    configure_theme(DEFAULT_THEME)
    yield
    configure_theme(DEFAULT_THEME)


def pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def slide_text(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def test_default_theme_has_no_logo_or_footer():
    assert not DEFAULT_THEME.footer_text
    assert DEFAULT_THEME.logo_dark is None
    assert DEFAULT_THEME.logo_red is None
    assert DEFAULT_THEME.logo_white is None

    prs = new_presentation()
    base = base_slide(prs, "Title", "Section", 1)
    chapter = chapter_slide(prs, 1, "Title", "Subtitle", "")
    assert not pictures(base)
    assert not pictures(chapter)
    assert "TEST BRAND" not in slide_text(base)


def test_custom_theme_preserves_logo_and_footer(tmp_path):
    from PIL import Image

    logo = tmp_path / "logo.png"
    Image.new("RGB", (12, 12), (20, 40, 60)).save(logo)
    custom = replace(
        DEFAULT_THEME,
        footer_text="TEST BRAND",
        logo_dark=logo,
        logo_red=logo,
        logo_white=logo,
    )
    assert custom.footer_text == "TEST BRAND"
    for logo_path in (custom.logo_dark, custom.logo_red, custom.logo_white):
        assert logo_path is not None and logo_path.is_file()

    configure_theme(custom)
    prs = new_presentation()
    base = base_slide(prs, "Title", "Section", 1)
    chapter = chapter_slide(prs, 1, "Title", "Subtitle", "")
    assert len(pictures(base)) == 1
    assert len(pictures(chapter)) == 1
    assert "TEST BRAND" in slide_text(base)


def test_multiline_citation_expands_upward_without_moving_bottom():
    from pptx.util import Inches

    prs = new_presentation()
    single = base_slide(prs, "Title", "Section", 1, "First reference")
    multiple = base_slide(prs, "Title", "Section", 2, "First reference\nSecond reference")
    single_citation = single.shapes[-1]
    multiple_citation = multiple.shapes[-1]

    assert single_citation.top == Inches(6.78)
    assert multiple_citation.top == Inches(6.58)
    assert single_citation.top + single_citation.height == Inches(7.02)
    assert multiple_citation.top + multiple_citation.height == Inches(7.02)
    assert len(multiple_citation.text_frame.paragraphs) == 2


def test_theme_replace_derives_an_independent_immutable_theme():
    accent = RGBColor(1, 2, 3)
    custom = replace(DEFAULT_THEME, accent=accent, footer_text="Group", logo_red=Path("logo.png"))
    assert custom.accent == accent
    assert custom.footer_text == "Group"
    assert custom.logo_red == Path("logo.png")
    assert DEFAULT_THEME.footer_text == ""
    assert DEFAULT_THEME.logo_red is None
    with pytest.raises(FrozenInstanceError):
        custom.footer_text = "Changed"


def test_generated_text_xml_has_script_typefaces_and_explicit_size(tmp_path):
    prs = new_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, "中文 English", 1, 1, 5, 0.5, size=23)
    rich_text(slide, [("汉", ""), ("2", "sup")], 1, 2, 5, 0.5, size=24)
    path = tmp_path / "fonts.pptx"
    prs.save(path)

    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("ppt/slides/slide1.xml"))
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    runs = {}
    for run in root.findall(".//a:r", ns):
        content = run.findtext("a:t", namespaces=ns)
        if content in {"中文", " English", "汉", "2"}:
            runs[content] = run.find("a:rPr", ns)

    for content, typeface in (
        ("中文", "Microsoft YaHei"),
        ("汉", "Microsoft YaHei"),
        (" English", "Arial"),
        ("2", "Arial"),
    ):
        properties = runs[content]
        assert properties is not None
        for tag in ("latin", "ea", "cs"):
            marker = properties.find(f"a:{tag}", ns)
            assert marker is not None
            assert marker.get("typeface") == typeface
    assert runs["中文"].get("sz") == "2300"
    assert runs[" English"].get("sz") == "2300"


def test_starter_builds_with_default_theme(tmp_path):
    source = ROOT / ".agents" / "skills" / "document" / "templates" / "starter" / "build_report.py"
    script = tmp_path / "build_report.py"
    shutil.copy2(source, script)
    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT)

    subprocess.run(
        [sys.executable, str(script)],
        cwd=tmp_path,
        env=env,
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )

    prs = Presentation(str(tmp_path / "starter.pptx"))
    assert len(prs.slides) == 2
    for slide in prs.slides:
        assert not pictures(slide)
        assert "PEKING UNIVERSITY" not in slide_text(slide)
