from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest

pytest.importorskip("pptx")

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from scrinium.presentation.editable import compare_pptx, prepare_editable


def _make_pptx(path: Path, slide_texts=("Before",)) -> Path:
    prs = Presentation()
    for text in slide_texts:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.name = f"Text: {text}"
        shape.text = text
    prs.save(str(path))
    return path


def _rewrite_parts(path: Path, updates: dict[str, bytes], timestamp=None) -> None:
    temporary = path.with_suffix(".tmp")
    with ZipFile(path) as source, ZipFile(temporary, "w") as target:
        for info in source.infolist():
            if timestamp is not None:
                info.date_time = timestamp
            target.writestr(info, updates.get(info.filename, source.read(info)))
    temporary.replace(path)


def _edited_xml(path: Path, part: str) -> etree._Element:
    with ZipFile(path) as archive:
        return etree.fromstring(archive.read(part))


def _xml_bytes(root: etree._Element, pretty_print=False) -> bytes:
    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=True,
        pretty_print=pretty_print,
    )


def test_prepare_editable_does_not_overwrite_by_default(tmp_path):
    source = _make_pptx(tmp_path / "source.pptx")
    target = tmp_path / "source_editable.pptx"
    target.write_bytes(b"keep me")

    with pytest.raises(FileExistsError):
        prepare_editable(source, target)

    assert target.read_bytes() == b"keep me"


def test_prepare_editable_force_overwrites_and_copy_has_no_diff(tmp_path):
    source = _make_pptx(tmp_path / "source.pptx")
    target = tmp_path / "source_editable.pptx"
    target.write_bytes(b"stale")

    assert prepare_editable(source, target, force=True) == target

    assert target.read_bytes() == source.read_bytes()
    assert compare_pptx(source, target)["has_changes"] is False


def test_compare_pptx_detects_text_and_geometry_changes(tmp_path):
    baseline = _make_pptx(tmp_path / "baseline.pptx")
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))
    shape = prs.slides[0].shapes[0]
    shape.text = "After"
    original_left = shape.left
    original_width = shape.width
    shape.left = original_left + Inches(0.5)
    shape.width = original_width + Inches(1)
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)
    slide = result["slides_modified"][0]

    assert result["summary"]["slides_modified"] == 1
    assert slide["text_changes"][0]["before"] == "Before"
    assert slide["text_changes"][0]["after"] == "After"
    assert slide["text_changes"][0]["match"] == "shape_id"
    geometry = slide["geometry_changes"][0]
    assert geometry["after"]["left"] - geometry["before"]["left"] == Inches(0.5)
    assert geometry["after"]["width"] - geometry["before"]["width"] == Inches(1)


def test_compare_pptx_detects_shape_and_slide_additions_and_removals(tmp_path):
    baseline = _make_pptx(tmp_path / "baseline.pptx", ("First", "Second"))
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))

    first_slide = prs.slides[0]
    added_shape = first_slide.shapes.add_textbox(Inches(2), Inches(2), Inches(3), Inches(1))
    added_shape.name = "Added shape"
    added_shape.text = "Added"
    removed_element = first_slide.shapes[0]._element
    removed_element.getparent().remove(removed_element)

    new_slide = prs.slides.add_slide(prs.slide_layouts[6])
    new_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "New slide"
    slide_ids = prs.slides._sldIdLst
    removed_slide_id = slide_ids[1]
    prs.part.drop_rel(removed_slide_id.rId)
    slide_ids.remove(removed_slide_id)
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)

    assert len(result["slides_added"]) == 1
    assert len(result["slides_removed"]) == 1
    assert result["slides_added"][0]["index"] == 2
    assert result["slides_removed"][0]["index"] == 2
    slide = result["slides_modified"][0]
    assert [shape["name"] for shape in slide["shapes_added"]] == ["Added shape"]
    assert [shape["name"] for shape in slide["shapes_removed"]] == ["Text: First"]


def test_compare_pptx_detects_image_replacement_and_notes_change(tmp_path):
    from PIL import Image

    before_image = tmp_path / "before.png"
    after_image = tmp_path / "after.png"
    Image.new("RGB", (8, 8), (255, 0, 0)).save(before_image)
    Image.new("RGB", (8, 8), (0, 0, 255)).save(after_image)

    baseline = tmp_path / "baseline.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(before_image), Inches(1), Inches(1), width=Inches(1), height=Inches(1))
    prs.save(str(baseline))
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)

    prs = Presentation(str(edited))
    slide = prs.slides[0]
    picture = slide.shapes[0]
    picture_element = picture._element
    picture_element.getparent().remove(picture_element)
    slide.shapes.add_picture(str(after_image), Inches(1), Inches(1), width=Inches(1), height=Inches(1))
    slide.notes_slide.notes_text_frame.text = "Revised notes"
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)
    slide_diff = result["slides_modified"][0]

    assert len(slide_diff["image_changes"]) == 1
    image_change = slide_diff["image_changes"][0]
    assert image_change["before"]["sha256"] != image_change["after"]["sha256"]
    assert slide_diff["notes_change"] == {"before": "", "after": "Revised notes"}


@pytest.mark.parametrize("font_change", ["name", "size", "bold", "color"])
def test_compare_pptx_detects_font_only_changes(tmp_path, font_change):
    baseline = _make_pptx(tmp_path / "baseline.pptx")
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))
    font = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font
    if font_change == "name":
        font.name = "Arial"
    elif font_change == "size":
        font.size = Pt(22)
    elif font_change == "bold":
        font.bold = True
    else:
        font.color.rgb = RGBColor(1, 2, 3)
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is True
    assert result["slides_modified"][0]["text_changes"] == []
    assert len(result["property_changes"]) == 1
    change = result["property_changes"][0]
    assert "font" in change["aspects"]
    assert change["name"] == "Text: Before"
    assert change["baseline_slide"] == change["edited_slide"] == 1
    assert change["before"]["signature"] != change["after"]["signature"]
    json.dumps(result)


def test_compare_pptx_detects_slide_reorder_only(tmp_path):
    baseline = _make_pptx(tmp_path / "baseline.pptx", ("First", "Second", "Third"))
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))
    original_ids = [slide.slide_id for slide in prs.slides]
    slide_ids = prs.slides._sldIdLst
    last_slide_id = list(slide_ids)[-1]
    slide_ids.remove(last_slide_id)
    slide_ids.insert(0, last_slide_id)
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is True
    assert result["slides_added"] == result["slides_removed"] == result["slides_modified"] == []
    assert result["summary"]["slide_order_changes"] == 1
    order_change = result["slide_order_change"]
    assert [slide["slide_id"] for slide in order_change["before"]] == original_ids
    assert [slide["slide_id"] for slide in order_change["after"]] == [
        original_ids[2],
        original_ids[0],
        original_ids[1],
    ]


@pytest.mark.parametrize("operation", ["add", "remove"])
def test_compare_pptx_slide_add_remove_does_not_reorder_common_slides(tmp_path, operation):
    baseline = _make_pptx(tmp_path / "baseline.pptx", ("First", "Second", "Third"))
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))
    slide_ids = prs.slides._sldIdLst
    if operation == "add":
        new_slide = prs.slides.add_slide(prs.slide_layouts[6])
        new_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "New"
        new_slide_id = list(slide_ids)[-1]
        slide_ids.remove(new_slide_id)
        slide_ids.insert(1, new_slide_id)
    else:
        removed_slide_id = slide_ids[1]
        prs.part.drop_rel(removed_slide_id.rId)
        slide_ids.remove(removed_slide_id)
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is True
    assert result["slide_order_change"] is None
    assert result["summary"]["slide_order_changes"] == 0
    assert len(result["slides_added"]) == int(operation == "add")
    assert len(result["slides_removed"]) == int(operation == "remove")


def test_compare_pptx_detects_picture_crop_only(tmp_path):
    from PIL import Image

    image_path = tmp_path / "image.png"
    Image.new("RGB", (8, 8), (255, 0, 0)).save(image_path)
    baseline = tmp_path / "baseline.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), width=Inches(1), height=Inches(1))
    prs.save(str(baseline))
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))
    prs.slides[0].shapes[0].crop_left = 0.25
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is True
    slide_diff = result["slides_modified"][0]
    assert slide_diff["image_changes"] == []
    assert slide_diff["geometry_changes"] == []
    assert len(result["property_changes"]) == 1
    assert "crop" in result["property_changes"][0]["aspects"]


def test_compare_pptx_detects_rotation_only(tmp_path):
    baseline = _make_pptx(tmp_path / "baseline.pptx")
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))
    prs.slides[0].shapes[0].rotation = 30
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is True
    assert result["slides_modified"][0]["geometry_changes"] == []
    assert len(result["property_changes"]) == 1
    assert "rotation" in result["property_changes"][0]["aspects"]


@pytest.mark.parametrize(
    ("part", "attribute", "value"),
    [
        ("ppt/slideMasters/slideMaster1.xml", "preserve", "0"),
        ("ppt/theme/theme1.xml", "name", "Edited Theme"),
    ],
)
def test_compare_pptx_detects_master_and_theme_xml_changes(tmp_path, part, attribute, value):
    baseline = _make_pptx(tmp_path / "baseline.pptx")
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    root = _edited_xml(baseline, part)
    root.set(attribute, value)
    _rewrite_parts(edited, {part: _xml_bytes(root)})

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is True
    assert result["slides_modified"] == []
    changes = [change for change in result["package_changes"]["modified"] if change["part"] == part]
    assert len(changes) == 1
    assert changes[0]["comparison"] == "canonical_xml"
    assert changes[0]["before_sha256"] != changes[0]["after_sha256"]


def test_compare_pptx_detects_media_binary_change_in_package(tmp_path):
    from PIL import Image

    before_image = tmp_path / "before.png"
    Image.new("RGB", (8, 8), (255, 0, 0)).save(before_image)
    baseline = tmp_path / "baseline.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(before_image), Inches(1), Inches(1))
    prs.save(str(baseline))
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    _rewrite_parts(edited, {"ppt/media/image1.png": b"arbitrary media bytes"})

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is True
    assert len(result["slides_modified"][0]["image_changes"]) == 1
    changes = [change for change in result["package_changes"]["modified"] if change["part"] == "ppt/media/image1.png"]
    assert len(changes) == 1
    assert changes[0]["kind"] == "binary"
    assert changes[0]["comparison"] == "content_hash"
    assert changes[0]["before_sha256"] != changes[0]["after_sha256"]


@pytest.mark.parametrize(
    ("part", "tag", "value"),
    [
        (
            "docProps/core.xml",
            "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastModifiedBy",
            "Reviewer",
        ),
        (
            "docProps/app.xml",
            "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}TotalTime",
            "42",
        ),
    ],
)
def test_compare_pptx_separates_pure_metadata_changes(tmp_path, part, tag, value):
    baseline = _make_pptx(tmp_path / "baseline.pptx")
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    root = _edited_xml(baseline, part)
    element = root.find(tag)
    if element is None:
        element = etree.SubElement(root, tag)
    element.text = value
    _rewrite_parts(edited, {part: _xml_bytes(root)})

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is False
    assert result["has_metadata_changes"] is True
    assert result["package_changes"] == {"added": [], "removed": [], "modified": []}
    assert [change["part"] for change in result["metadata_changes"]["modified"]] == [part]
    assert result["summary"]["package_changes"] == 0
    assert result["summary"]["metadata_changes"] == 1


def test_compare_pptx_ignores_xml_formatting_and_zip_timestamps(tmp_path):
    baseline = _make_pptx(tmp_path / "baseline.pptx")
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    part = "ppt/theme/theme1.xml"
    root = _edited_xml(baseline, part)
    for element in root.iter():
        attributes = list(element.attrib.items())
        element.attrib.clear()
        element.attrib.update(reversed(attributes))
    _rewrite_parts(
        edited,
        {part: _xml_bytes(root, pretty_print=True)},
        timestamp=(2030, 1, 1, 0, 0, 0),
    )

    result = compare_pptx(baseline, edited)

    assert result["has_changes"] is False
    assert result["has_metadata_changes"] is False
    assert result["summary"]["package_changes"] == 0


def test_compare_pptx_covers_text_inside_group_shapes(tmp_path):
    baseline = tmp_path / "baseline.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    textbox = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.name = "Nested text"
    textbox.text = "Before"
    prs.save(str(baseline))
    edited = tmp_path / "edited.pptx"
    prepare_editable(baseline, edited)
    prs = Presentation(str(edited))
    nested_textbox = prs.slides[0].shapes[0].shapes[0]
    nested_textbox.text_frame.paragraphs[0].runs[0].text = "After"
    prs.save(str(edited))

    result = compare_pptx(baseline, edited)

    text_changes = result["slides_modified"][0]["text_changes"]
    assert len(text_changes) == 1
    assert text_changes[0]["name"] == "Nested text"
    assert text_changes[0]["edited_path"] == [1, 1]
    assert text_changes[0]["before"] == "Before"
    assert text_changes[0]["after"] == "After"
