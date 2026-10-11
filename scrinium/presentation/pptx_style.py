from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from scrinium.presentation.theme import DEFAULT_THEME

SLIDE_W = 13.333
SLIDE_H = 7.5
CURRENT_THEME = DEFAULT_THEME
FONT_CN = DEFAULT_THEME.font_cn
FONT_EN = DEFAULT_THEME.font_en

DARK = RGBColor(39, 43, 48)
MID = RGBColor(96, 103, 112)
LIGHT = RGBColor(221, 225, 229)
PALE = RGBColor(246, 242, 242)
PALE_BLUE = RGBColor(240, 245, 249)
BLUE = RGBColor(54, 104, 143)
GREEN = RGBColor(55, 125, 103)
AMBER = RGBColor(176, 117, 28)
WHITE = RGBColor(255, 255, 255)


def configure_theme(theme):
    global CURRENT_THEME, FONT_CN, FONT_EN
    CURRENT_THEME = theme
    FONT_CN = theme.font_cn
    FONT_EN = theme.font_en


def I(value):  # noqa: E743
    return Inches(value)


def set_typefaces(run, font):
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            rpr.append(element)
        element.set("typeface", font)


def _is_cjk(char):
    code = ord(char)
    return 0x3400 <= code <= 0x4DBF or 0x4E00 <= code <= 0x9FFF or 0xF900 <= code <= 0xFAFF


def add_scripted_runs(paragraph, content, font):
    runs = []

    def add_one(segment, segment_font):
        run = paragraph.add_run()
        run.text = segment
        set_typefaces(run, segment_font)
        runs.append(run)

    current = []
    current_font = None
    for char in content:
        char_font = font if _is_cjk(char) else FONT_EN
        if current_font is None:
            current_font = char_font
        if char_font != current_font and current:
            add_one("".join(current), current_font)
            current = []
        current_font = char_font
        current.append(char)
    if current:
        add_one("".join(current), current_font)
    return runs


def set_theme_east_asian_font(path, typeface, latin=None):
    latin = FONT_EN if latin is None else latin
    tmp_path = path.with_suffix(".tmp")
    with ZipFile(path) as src, ZipFile(tmp_path, "w", ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.startswith("ppt/theme/theme") and info.filename.endswith(".xml"):
                root = etree.fromstring(data)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for element in root.xpath('.//a:ea | .//a:font[@script="Hans"]', namespaces=ns):
                    element.set("typeface", typeface)
                for element in root.xpath(".//a:latin", namespaces=ns):
                    element.set("typeface", latin)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
            dst.writestr(info, data)
    tmp_path.replace(path)


def new_presentation():
    prs = Presentation()
    prs.slide_width = I(SLIDE_W)
    prs.slide_height = I(SLIDE_H)
    return prs


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, I(x), I(y), I(w), I(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.8)
    return shape


def line(slide, x1, y1, x2, y2, color=LIGHT, width=1.0):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def text(
    slide,
    content,
    x,
    y,
    w,
    h,
    size=20,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font=None,
    margin=0,
    line_spacing=1.0,
):
    font = FONT_CN if font is None else font
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(margin)
    tf.vertical_anchor = valign
    lines = str(content).split("\n")
    for idx, content_line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = line_spacing
        for run in add_scripted_runs(p, content_line, font):
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def rich_text(
    slide,
    parts,
    x,
    y,
    w,
    h,
    size=20,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font=None,
    margin=0,
):
    font = FONT_CN if font is None else font
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    for content, kind in parts:
        for run in add_scripted_runs(p, content, font):
            run.font.size = Pt(size if kind == "" else size * 0.68)
            run.font.bold = bold
            run.font.color.rgb = color
            if kind:
                run._r.get_or_add_rPr().set("baseline", "25000" if kind == "sup" else "-25000")
    return box


def _existing_logo(path):
    if path is None:
        return None
    path = Path(path)
    return path if path.exists() else None


def base_slide(prs, title, section, number, citation=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = CURRENT_THEME.accent
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    text(slide, section, 0.58, 0.24, 7.4, 0.28, 11, accent, True)
    text(slide, title, 0.58, 0.59, 12.1, 0.58, 28, DARK, True)
    rect(slide, 0.58, 1.28, 0.78, 0.035, accent)
    line(slide, 1.36, 1.297, 12.75, 1.297, LIGHT, 0.8)
    line(slide, 0.45, 7.06, 12.88, 7.06, LIGHT, 0.7)
    logo_dark = _existing_logo(CURRENT_THEME.logo_dark)
    if logo_dark is not None:
        slide.shapes.add_picture(str(logo_dark), I(0.47), I(7.15), width=I(0.23), height=I(0.23))
    if CURRENT_THEME.footer_text:
        text(slide, CURRENT_THEME.footer_text, 0.77, 7.17, 2.0, 0.18, 8.5, MID, font=FONT_EN)
    text(slide, f"{len(prs.slides):02d}", 12.2, 7.14, 0.55, 0.2, 10, MID, align=PP_ALIGN.RIGHT, font=FONT_EN)
    if citation:
        extra_lines = str(citation).count("\n")
        text(slide, citation, 0.58, 6.78 - 0.20 * extra_lines, 11.9, 0.24 + 0.20 * extra_lines, 10, MID)
    return slide


def image_contain(slide, path, x, y, w, h, caption=None, caption_size=18, border=False, background=WHITE):
    path = Path(path)
    if caption:
        caption_h = 0.38
        image_h = h - caption_h
    else:
        caption_h = 0
        image_h = h
    rect(slide, x, y, w, image_h, background, LIGHT if border else None)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, image_h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2
    py = y + (image_h - ph) / 2
    slide.shapes.add_picture(str(path), I(px), I(py), width=I(pw), height=I(ph))
    if caption:
        text(
            slide,
            caption,
            x,
            y + image_h + 0.04,
            w,
            caption_h - 0.04,
            caption_size,
            MID,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )


def label(slide, content, x, y, w, color=None, fill=PALE, size=13):
    color = CURRENT_THEME.accent if color is None else color
    rect(slide, x, y, w, 0.35, fill, None, radius=True)
    text(
        slide,
        content,
        x + 0.06,
        y + 0.035,
        w - 0.12,
        0.26,
        size,
        color,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def card(slide, title, body, x, y, w, h, accent=None, fill=WHITE, title_size=19, body_size=18):
    accent = CURRENT_THEME.accent if accent is None else accent
    body_h = h - 0.88
    line_count = body.count("\n") + 1
    needed_h = body_size / 72 * (1.12 * line_count - 0.12)
    if needed_h > body_h:
        raise ValueError(f"Card '{title}' needs {needed_h:.2f} in for body, only {body_h:.2f} in available")
    rect(slide, x, y, w, h, fill, LIGHT, radius=True)
    rect(slide, x, y, 0.055, h, accent)
    text(slide, title, x + 0.22, y + 0.18, w - 0.38, 0.42, title_size, DARK, True)
    text(slide, body, x + 0.22, y + 0.70, w - 0.38, body_h, body_size, MID, valign=MSO_ANCHOR.MIDDLE, line_spacing=1.12)


def right_arrow(slide, x, y, w=0.42, h=0.28, color=None):
    color = CURRENT_THEME.accent if color is None else color
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, I(x), I(y), I(w), I(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def circle(slide, x, y, d, fill=None, line_color=None, width=0.8):
    fill = CURRENT_THEME.accent if fill is None else fill
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(width)
    return shape


def add_note(slide, content, x=0.58, y=6.28, w=12.1, color=MID, size=15):
    text(slide, content, x, y, w, 0.32, size, color)


def chapter_slide(prs, number, title, subtitle, source, english=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    rect(slide, 0, 0, 3.65, SLIDE_H, CURRENT_THEME.accent)
    logo_white = _existing_logo(CURRENT_THEME.logo_white)
    if logo_white is not None:
        slide.shapes.add_picture(str(logo_white), I(0.62), I(0.52), width=I(0.56), height=I(0.56))
    text(slide, f"{number:02d}", 0.62, 1.48, 2.35, 0.72, 44, WHITE, True, font=FONT_EN)
    text(slide, "CHAPTER", 0.65, 2.28, 2.20, 0.28, 12, WHITE, font=FONT_EN)
    rect(slide, 0.65, 2.75, 0.72, 0.035, WHITE)
    subtitle_size = 21 if len(subtitle) > 6 else 24
    text(slide, subtitle, 0.65, 3.16, 2.75, 1.25, subtitle_size, WHITE, True, line_spacing=1.08)
    text(slide, title, 4.35, 2.55, 8.10, 1.05, 34, DARK, True, valign=MSO_ANCHOR.MIDDLE)
    line(slide, 4.35, 3.92, 12.35, 3.92, LIGHT, 0.9)
    if english:
        text(slide, english, 4.35, 4.20, 7.85, 0.35, 13, MID, font=FONT_EN)
    text(slide, f"{len(prs.slides):02d}", 12.2, 7.14, 0.55, 0.2, 10, MID, align=PP_ALIGN.RIGHT, font=FONT_EN)
    return slide


def horizontal_bars(
    slide,
    title,
    items,
    x,
    y,
    w,
    h,
    max_value=None,
    label_w=1.65,
    value_w=1.0,
    title_size=18,
    label_size=18,
    value_size=18,
    colors=None,
    value_format=None,
):
    if not items:
        return
    text(slide, title, x, y, w, 0.32, title_size, DARK, True)
    max_value = max_value or max(value for _, value in items)
    row_h = (h - 0.42) / len(items)
    bar_x = x + label_w
    bar_w = w - label_w - value_w
    for i, (label, value) in enumerate(items):
        cy = y + 0.42 + i * row_h
        color = colors[i] if colors else CURRENT_THEME.accent
        text(slide, label, x, cy + 0.03, label_w - 0.06, row_h - 0.05, label_size, MID, valign=MSO_ANCHOR.MIDDLE)
        rect(slide, bar_x, cy + 0.10, max(0.015, bar_w * value / max_value), max(0.10, row_h - 0.20), color)
        shown = value_format(value) if value_format else f"{value:g}"
        text(
            slide,
            shown,
            bar_x + bar_w + 0.06,
            cy + 0.03,
            value_w - 0.06,
            row_h - 0.05,
            value_size,
            color,
            True,
            align=PP_ALIGN.RIGHT,
            valign=MSO_ANCHOR.MIDDLE,
            font=FONT_EN,
        )
