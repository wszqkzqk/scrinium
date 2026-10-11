"""Create editable PPTX copies and compare presentation revisions."""

from __future__ import annotations

import hashlib
import shutil
from collections import Counter
from pathlib import Path
from typing import Any


def _require_pptx(path: str | Path, label: str) -> Path:
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"{label}不存在: {file_path}")
    if not file_path.is_file():
        raise ValueError(f"{label}不是文件: {file_path}")
    if file_path.suffix.lower() != ".pptx":
        raise ValueError(f"{label}不是 PPTX 文件: {file_path}")
    return file_path


def prepare_editable(source: str | Path, target: str | Path, force: bool = False) -> Path:
    """Copy a PPTX for manual editing without implicitly overwriting the target."""
    source_path = _require_pptx(source, "源文件")
    target_path = Path(target)
    if target_path.suffix.lower() != ".pptx":
        raise ValueError(f"输出路径不是 PPTX 文件: {target_path}")
    if source_path.resolve() == target_path.resolve():
        raise ValueError("可编辑副本不能覆盖源文件")
    if target_path.exists():
        if target_path.is_dir():
            raise ValueError(f"输出路径不是文件: {target_path}")
        if not force:
            raise FileExistsError(f"可编辑副本已存在: {target_path}（如需覆盖请使用 force=True）")
    if not target_path.parent.is_dir():
        raise FileNotFoundError(f"输出目录不存在: {target_path.parent}")
    shutil.copy2(source_path, target_path)
    return target_path


def _unique_keys(items, indices, key_func):
    values = {index: key_func(items[index]) for index in indices}
    counts = Counter(value for value in values.values() if value is not None)
    return {value: index for index, value in values.items() if value is not None and counts[value] == 1}


def _match_items(baseline, edited, key_func, name_func=None):
    baseline_unmatched = set(range(len(baseline)))
    edited_unmatched = set(range(len(edited)))
    baseline_ids = {index: key_func(item) for index, item in enumerate(baseline)}
    edited_ids = {index: key_func(item) for index, item in enumerate(edited)}
    baseline_id_counts = Counter(value for value in baseline_ids.values() if value is not None)
    edited_id_counts = Counter(value for value in edited_ids.values() if value is not None)
    pairs = []

    def pair_unique(value_func, match_kind):
        baseline_keys = _unique_keys(baseline, baseline_unmatched, value_func)
        edited_keys = _unique_keys(edited, edited_unmatched, value_func)
        shared = baseline_keys.keys() & edited_keys.keys()
        matched = [(baseline_keys[value], edited_keys[value], match_kind) for value in sorted(shared, key=str)]
        for baseline_index, edited_index, _ in matched:
            baseline_unmatched.remove(baseline_index)
            edited_unmatched.remove(edited_index)
        pairs.extend(matched)

    pair_unique(key_func, "shape_id" if name_func is not None else "slide_id")
    if name_func is not None:
        pair_unique(name_func, "name")

    for index in sorted(baseline_unmatched & edited_unmatched):
        baseline_id = baseline_ids[index]
        edited_id = edited_ids[index]
        baseline_reliable = baseline_id is not None and baseline_id_counts[baseline_id] == 1
        edited_reliable = edited_id is not None and edited_id_counts[edited_id] == 1
        if baseline_reliable and edited_reliable:
            continue
        pairs.append((index, index, "index"))
        baseline_unmatched.remove(index)
        edited_unmatched.remove(index)

    return (
        sorted(pairs, key=lambda pair: pair[1]),
        sorted(baseline_unmatched),
        sorted(edited_unmatched),
    )


def _slide_ref(slide, index):
    return {"slide_id": slide.slide_id, "index": index + 1}


def _shape_ref(shape, path):
    indices = [index + 1 for index in path]
    return {
        "shape_id": shape.shape_id,
        "name": shape.name,
        "index": indices[-1],
        "path": indices,
        "shape_type": str(shape.shape_type),
    }


def _change_ref(baseline_entry, edited_entry, match_kind):
    baseline_shape, baseline_path = baseline_entry
    edited_shape, edited_path = edited_entry
    return {
        "shape_id": edited_shape.shape_id,
        "name": edited_shape.name,
        "baseline_index": baseline_path[-1] + 1,
        "edited_index": edited_path[-1] + 1,
        "baseline_path": [index + 1 for index in baseline_path],
        "edited_path": [index + 1 for index in edited_path],
        "match": match_kind,
    }


def _flatten_shapes(shapes, prefix=()):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    entries = []
    for index, shape in enumerate(shapes):
        path = (*prefix, index)
        entries.append((shape, path))
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            entries.extend(_flatten_shapes(shape.shapes, path))
    return entries


def _shape_text(shape):
    if shape.has_text_frame:
        return shape.text
    if shape.has_table:
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells)
    return None


def _geometry(shape):
    return {key: int(getattr(shape, key) or 0) for key in ("left", "top", "width", "height")}


def _image_info(shape):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return None
    relationship_id = shape._element.blip_rId
    if relationship_id is None:
        return None
    image_part = shape.part.related_part(relationship_id)
    return {
        "sha256": hashlib.sha256(image_part.blob).hexdigest(),
        "content_type": image_part.content_type,
    }


def _notes_text(slide):
    if not slide.has_notes_slide:
        return ""
    text_frame = slide.notes_slide.notes_text_frame
    return text_frame.text if text_frame is not None else ""


def _canonical_xml(value, ignore_text=False):
    from lxml import etree

    if isinstance(value, bytes):
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        root = etree.fromstring(value, parser)
    else:
        root = etree.fromstring(etree.tostring(value))

    for element in root.iter():
        if ignore_text:
            element.text = None
            element.tail = None
            continue
        children = list(element)
        if not children:
            continue
        direct_text = [element.text] + [child.tail for child in children]
        if any(text and text.strip() for text in direct_text):
            continue
        element.text = None
        for child in children:
            child.tail = None

    return etree.tostring(root, method="c14n", with_comments=True)


def _xml_signature(elements):
    digest = hashlib.sha256()
    for element in elements:
        digest.update(_canonical_xml(element, ignore_text=True))
        digest.update(b"\0")
    return digest.hexdigest()


def _property_root(shape):
    from copy import deepcopy

    from pptx.enum.shapes import MSO_SHAPE_TYPE

    root = deepcopy(shape._element)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        namespace = "http://schemas.openxmlformats.org/presentationml/2006/main"
        child_tags = {
            f"{{{namespace}}}{name}" for name in ("sp", "grpSp", "graphicFrame", "cxnSp", "pic", "contentPart")
        }
        for child in list(root):
            if child.tag in child_tags:
                root.remove(child)
    return root


def _shape_property_signatures(shape):
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    aspect_tags = {
        "font": {"rPr", "defRPr", "endParaRPr", "fontRef"},
        "paragraph": {"pPr"},
        "crop": {"srcRect"},
        "shape": {"spPr", "grpSpPr", "style"},
    }
    root = _property_root(shape)
    signatures = {"xml": _xml_signature([root])}
    for aspect, names in aspect_tags.items():
        tags = {f"{{{drawing_ns}}}{name}" for name in names}
        signatures[aspect] = _xml_signature(element for element in root.iter() if element.tag in tags)
    rotation_values = [
        element.get("rot")
        for element in root.iter()
        if element.tag == f"{{{drawing_ns}}}xfrm" and element.get("rot") is not None
    ]
    signatures["rotation"] = hashlib.sha256("\0".join(rotation_values).encode()).hexdigest()
    return signatures


def _compare_slide(baseline_slide, baseline_index, edited_slide, edited_index):
    baseline_shapes = _flatten_shapes(baseline_slide.shapes)
    edited_shapes = _flatten_shapes(edited_slide.shapes)
    pairs, removed, added = _match_items(
        baseline_shapes,
        edited_shapes,
        lambda entry: entry[0].shape_id,
        lambda entry: entry[0].name,
    )
    result = {
        "slide_id": edited_slide.slide_id,
        "baseline_slide": baseline_index + 1,
        "edited_slide": edited_index + 1,
        "shapes_added": [_shape_ref(*edited_shapes[index]) for index in added],
        "shapes_removed": [_shape_ref(*baseline_shapes[index]) for index in removed],
        "text_changes": [],
        "geometry_changes": [],
        "image_changes": [],
        "property_changes": [],
        "notes_change": None,
    }

    for baseline_shape_index, edited_shape_index, match_kind in pairs:
        baseline_entry = baseline_shapes[baseline_shape_index]
        edited_entry = edited_shapes[edited_shape_index]
        baseline_shape = baseline_entry[0]
        edited_shape = edited_entry[0]
        ref = _change_ref(baseline_entry, edited_entry, match_kind)
        baseline_text = _shape_text(baseline_shape)
        edited_text = _shape_text(edited_shape)
        if baseline_text != edited_text:
            result["text_changes"].append({**ref, "before": baseline_text, "after": edited_text})

        baseline_geometry = _geometry(baseline_shape)
        edited_geometry = _geometry(edited_shape)
        if baseline_geometry != edited_geometry:
            result["geometry_changes"].append({**ref, "before": baseline_geometry, "after": edited_geometry})

        baseline_image = _image_info(baseline_shape)
        edited_image = _image_info(edited_shape)
        if baseline_image != edited_image:
            result["image_changes"].append({**ref, "before": baseline_image, "after": edited_image})

        baseline_properties = _shape_property_signatures(baseline_shape)
        edited_properties = _shape_property_signatures(edited_shape)
        if baseline_properties["xml"] != edited_properties["xml"]:
            aspects = [
                aspect
                for aspect in ("font", "paragraph", "crop", "rotation", "shape")
                if baseline_properties[aspect] != edited_properties[aspect]
            ]
            result["property_changes"].append(
                {
                    **ref,
                    "aspects": aspects or ["other"],
                    "before": {"signature": baseline_properties["xml"]},
                    "after": {"signature": edited_properties["xml"]},
                }
            )

    baseline_notes = _notes_text(baseline_slide)
    edited_notes = _notes_text(edited_slide)
    if baseline_notes != edited_notes:
        result["notes_change"] = {"before": baseline_notes, "after": edited_notes}

    change_keys = (
        "shapes_added",
        "shapes_removed",
        "text_changes",
        "geometry_changes",
        "image_changes",
        "property_changes",
        "notes_change",
    )
    if any(result[key] for key in change_keys):
        return result
    return None


def _slide_order_change(baseline_slides, edited_slides, pairs):
    baseline_order = sorted(pairs, key=lambda pair: pair[0])
    edited_order = sorted(pairs, key=lambda pair: pair[1])
    if [(pair[0], pair[1]) for pair in baseline_order] == [(pair[0], pair[1]) for pair in edited_order]:
        return None
    return {
        "before": [_slide_ref(baseline_slides[pair[0]], pair[0]) for pair in baseline_order],
        "after": [_slide_ref(edited_slides[pair[1]], pair[1]) for pair in edited_order],
    }


def _package_parts(path):
    from zipfile import ZipFile

    with ZipFile(path) as archive:
        return {info.filename: archive.read(info) for info in archive.infolist() if not info.is_dir()}


def _part_fingerprint(data):
    from lxml import etree

    sample = data[3:] if data.startswith(b"\xef\xbb\xbf") else data
    kind = "binary"
    normalized = data
    if sample.lstrip().startswith(b"<"):
        try:
            normalized = _canonical_xml(data)
            kind = "xml"
        except (etree.LxmlError, ValueError):
            pass
    normalized_digest = hashlib.sha256()
    normalized_digest.update(kind.encode())
    normalized_digest.update(b"\0")
    normalized_digest.update(normalized)
    return {
        "kind": kind,
        "sha256": hashlib.sha256(data).hexdigest(),
        "normalized_sha256": normalized_digest.hexdigest(),
    }


def _is_metadata_part(part):
    return part in {"docProps/core.xml", "docProps/app.xml"} or part.startswith("docProps/thumbnail.")


def _compare_packages(baseline_path, edited_path):
    baseline_parts = _package_parts(baseline_path)
    edited_parts = _package_parts(edited_path)
    baseline_info = {part: _part_fingerprint(data) for part, data in baseline_parts.items()}
    edited_info = {part: _part_fingerprint(data) for part, data in edited_parts.items()}
    package_changes = {"added": [], "removed": [], "modified": []}
    metadata_changes = {"added": [], "removed": [], "modified": []}

    for part in sorted(baseline_parts.keys() | edited_parts.keys()):
        changes = metadata_changes if _is_metadata_part(part) else package_changes
        if part not in baseline_parts:
            info = edited_info[part]
            changes["added"].append({"part": part, "change": "added", "kind": info["kind"], "sha256": info["sha256"]})
            continue
        if part not in edited_parts:
            info = baseline_info[part]
            changes["removed"].append(
                {"part": part, "change": "removed", "kind": info["kind"], "sha256": info["sha256"]}
            )
            continue
        before = baseline_info[part]
        after = edited_info[part]
        if before["normalized_sha256"] == after["normalized_sha256"]:
            continue
        changes["modified"].append(
            {
                "part": part,
                "change": "modified",
                "kind": after["kind"],
                "comparison": "canonical_xml" if before["kind"] == after["kind"] == "xml" else "content_hash",
                "before_sha256": before["sha256"],
                "after_sha256": after["sha256"],
            }
        )
    return package_changes, metadata_changes


def _package_change_count(changes):
    return sum(len(changes[kind]) for kind in ("added", "removed", "modified"))


def compare_pptx(baseline: str | Path, edited: str | Path) -> dict[str, Any]:
    """Return a JSON-serializable semantic diff with an OPC package-level fallback."""
    baseline_path = _require_pptx(baseline, "基准文件")
    edited_path = _require_pptx(edited, "编辑文件")
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("python-pptx 未安装，请运行: pip install scrinium[office]") from None

    baseline_prs = Presentation(str(baseline_path))
    edited_prs = Presentation(str(edited_path))
    baseline_slides = list(baseline_prs.slides)
    edited_slides = list(edited_prs.slides)
    pairs, removed, added = _match_items(
        baseline_slides,
        edited_slides,
        lambda slide: slide.slide_id,
    )
    slide_order_change = _slide_order_change(baseline_slides, edited_slides, pairs)

    slides_modified = []
    for baseline_index, edited_index, _ in pairs:
        slide_diff = _compare_slide(
            baseline_slides[baseline_index],
            baseline_index,
            edited_slides[edited_index],
            edited_index,
        )
        if slide_diff is not None:
            slides_modified.append(slide_diff)

    slides_added = [_slide_ref(edited_slides[index], index) for index in added]
    slides_removed = [_slide_ref(baseline_slides[index], index) for index in removed]
    property_changes = [
        {
            "slide_id": slide["slide_id"],
            "baseline_slide": slide["baseline_slide"],
            "edited_slide": slide["edited_slide"],
            **change,
        }
        for slide in slides_modified
        for change in slide["property_changes"]
    ]
    package_changes, metadata_changes = _compare_packages(baseline_path, edited_path)
    package_change_count = _package_change_count(package_changes)
    metadata_change_count = _package_change_count(metadata_changes)
    summary = {
        "slides_added": len(slides_added),
        "slides_removed": len(slides_removed),
        "slides_modified": len(slides_modified),
        "slide_order_changes": int(slide_order_change is not None),
        "shapes_added": sum(len(slide["shapes_added"]) for slide in slides_modified),
        "shapes_removed": sum(len(slide["shapes_removed"]) for slide in slides_modified),
        "text_changes": sum(len(slide["text_changes"]) for slide in slides_modified),
        "geometry_changes": sum(len(slide["geometry_changes"]) for slide in slides_modified),
        "image_changes": sum(len(slide["image_changes"]) for slide in slides_modified),
        "property_changes": len(property_changes),
        "notes_changes": sum(slide["notes_change"] is not None for slide in slides_modified),
        "package_changes": package_change_count,
        "metadata_changes": metadata_change_count,
    }
    return {
        "baseline": str(baseline_path),
        "edited": str(edited_path),
        "has_changes": bool(
            slides_added or slides_removed or slides_modified or slide_order_change or package_change_count
        ),
        "has_metadata_changes": bool(metadata_change_count),
        "slide_count": {"baseline": len(baseline_slides), "edited": len(edited_slides)},
        "summary": summary,
        "slides_added": slides_added,
        "slides_removed": slides_removed,
        "slides_modified": slides_modified,
        "slide_order_change": slide_order_change,
        "property_changes": property_changes,
        "package_changes": package_changes,
        "metadata_changes": metadata_changes,
    }


__all__ = ["compare_pptx", "prepare_editable"]
